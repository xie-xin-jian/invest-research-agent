"""Markdown → Word(docx) 简易转换：支持 # 标题、- 列表、| 表格 |、普通段落。"""
from __future__ import annotations

import io

import docx


def markdown_to_docx(md: str, title: str = "投研资料") -> bytes:
    """把 Markdown 文本转成 docx 文件字节，返回可直接上传/保存的 bytes。"""
    d = docx.Document()
    d.add_heading(title, 0)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # 表格块：连续的 | 开头行
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            # 去掉 markdown 分隔行（|---|---|）
            rows = [r for r in rows if not all(set(c) <= set("-: ") and c for c in r)]
            if rows:
                table = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        table.rows[ri].cells[ci].text = cell
            continue

        if line.startswith("### "):
            d.add_heading(line[4:], 3)
        elif line.startswith("## "):
            d.add_heading(line[3:], 2)
        elif line.startswith("# "):
            d.add_heading(line[2:], 1)
        elif line.startswith("- "):
            d.add_paragraph(line[2:], style="List Bullet")
        elif line:
            d.add_paragraph(line)
        i += 1

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
